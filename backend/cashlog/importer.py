"""Read-only import of 记账单.xlsx into SQLite.

Responsibilities (P1):
  - Load the 汇率 sheet into exchange_rate.
  - Import 支出 (expense) and 入账 (income) into txn with a stable external_uid,
    a content hash (for later diffing), and a Decimal-exact CNY conversion.
  - Classify rows that cannot be imported as-is (anomalies) instead of silently
    dropping or corrupting them.
  - Produce a migration report: reconciliation diffs vs the workbook's own
    values, anomalies, stale-rate rows, and AA candidates.

The workbook is never modified here.
"""
from __future__ import annotations

import hashlib
import sqlite3
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl

from . import config, money, rates
from .seed import currency_id_map

EXPENSE_SHEET = "支出"
INCOME_SHEET = "入账"
RATE_SHEET = "汇率"
STATS_SHEET = "统计"

# Deterministic namespace for initial-import UIDs. uuid5 over a stable name keeps
# external_uid identical across re-runs of the migration on an unchanged workbook.
# Once P2 writes _cashlog_uid back into the xlsx, that column takes precedence.
CASHLOG_NS = uuid.uuid5(uuid.NAMESPACE_DNS, "cashlog.local")


def make_external_uid(sheet: str, source_row: int, row_hash: str) -> str:
    return uuid.uuid5(CASHLOG_NS, f"{sheet}:{source_row}:{row_hash}").hex


@dataclass
class Report:
    rates_loaded: int = 0
    fiscal_months_loaded: int = 0
    expense_imported: int = 0
    income_imported: int = 0
    anomalies: list = field(default_factory=list)      # rows we did NOT import
    reconcile_diffs: list = field(default_factory=list)  # imported but value mismatch
    no_rate: list = field(default_factory=list)          # imported, no rate available
    computed_rmb: list = field(default_factory=list)     # imported, workbook had no RMB
    stale_rate: list = field(default_factory=list)       # rate older than threshold
    aa_candidates: list = field(default_factory=list)    # content contains 'AA'

    def _count_by_sheet(self, items: list, sheet: str) -> int:
        return sum(1 for it in items if it.get("sheet") == sheet)

    def summary(self) -> dict:
        income_anomalies = self._count_by_sheet(self.anomalies, INCOME_SHEET)
        expense_anomalies = self._count_by_sheet(self.anomalies, EXPENSE_SHEET)
        return {
            "rates_loaded": self.rates_loaded,
            "fiscal_months_loaded": self.fiscal_months_loaded,
            "expense_imported": self.expense_imported,
            "income_imported": self.income_imported,
            # Every non-blank income row is accounted for: imported OR quarantined.
            "income_rows_accounted": self.income_imported + income_anomalies,
            "expense_anomalies": expense_anomalies,
            "income_anomalies": income_anomalies,
            "anomalies": len(self.anomalies),
            "reconcile_diffs": len(self.reconcile_diffs),
            "no_rate": len(self.no_rate),
            "computed_rmb": len(self.computed_rmb),
            "stale_rate": len(self.stale_rate),
            # Split so the count is not misread: structured-AA candidates live in
            # expenses; income only *mentions* AA in its source text.
            "aa_expense_candidates": self._count_by_sheet(self.aa_candidates, EXPENSE_SHEET),
            "aa_income_mentions": self._count_by_sheet(self.aa_candidates, INCOME_SHEET),
        }


def _iso(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return None


def _row_hash(direction: str, txn_date: str, amount_minor: int, currency_code: str,
              content: str, category: str, note: str) -> str:
    raw = f"{direction}|{txn_date}|{amount_minor}|{currency_code}|{content}|{category}|{note}"
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def load_rates(conn: sqlite3.Connection, wb) -> int:
    ccy = currency_id_map(conn)
    ws = wb[RATE_SHEET]
    count = 0
    for r in range(2, ws.max_row + 1):
        d = ws.cell(row=r, column=1).value
        cu = ws.cell(row=r, column=2).value
        rt = ws.cell(row=r, column=3).value
        iso = _iso(d)
        if iso is None or cu not in ccy or money.parse_decimal(rt) is None:
            continue
        conn.execute(
            """INSERT INTO exchange_rate (currency_id, rate_date, rate_to_base, source)
               VALUES (?, ?, ?, 'workbook')
               ON CONFLICT(currency_id, rate_date) DO UPDATE
                   SET rate_to_base = excluded.rate_to_base""",
            (ccy[cu]["id"], iso, money.rate_to_str(rt)),
        )
        count += 1
    conn.commit()
    return count


def load_fiscal_months(conn: sqlite3.Connection, wb, ccy: dict) -> int:
    """Import the 统计 sheet's manual inputs (财政月 + 信用卡金额, USD) into
    fiscal_month. The derived columns (总支出/换算/占比) are recomputed, not stored."""
    if STATS_SHEET not in wb.sheetnames:
        return 0
    usd = ccy.get("美元") or ccy.get("USD")
    ws = wb[STATS_SHEET]
    count = 0
    for r in range(2, ws.max_row + 1):
        iso = _iso(ws.cell(row=r, column=1).value)   # A: 财政月 (the 5th)
        if iso is None:
            continue
        cc = money.parse_decimal(ws.cell(row=r, column=5).value)  # E: 信用卡金额 (USD)
        minor = money.to_minor(cc if cc is not None else 0, usd["minor_unit"])
        conn.execute(
            """INSERT INTO fiscal_month (period_end, credit_card_amount_minor, credit_card_ccy_id)
               VALUES (?, ?, ?)
               ON CONFLICT(period_end) DO UPDATE SET
                   credit_card_amount_minor = excluded.credit_card_amount_minor,
                   credit_card_ccy_id = excluded.credit_card_ccy_id""",
            (iso, minor, usd["id"]),
        )
        count += 1
    conn.commit()
    return count


class _CategoryCache:
    def __init__(self, conn: sqlite3.Connection):
        self.conn = conn
        self.cache: dict = {}
        for row in conn.execute("SELECT id, name, kind FROM category"):
            self.cache[(row["name"], row["kind"])] = row["id"]

    def get_or_create(self, name: Optional[str], kind: str) -> Optional[int]:
        if name is None or str(name).strip() == "":
            return None
        name = str(name).strip()
        key = (name, kind)
        if key in self.cache:
            return self.cache[key]
        cur = self.conn.execute(
            "INSERT INTO category (name, kind) VALUES (?, ?)", (name, kind)
        )
        self.cache[key] = cur.lastrowid
        return cur.lastrowid


def _import_sheet(conn: sqlite3.Connection, wb, sheet_name: str, direction: str,
                  ccy: dict, cats: _CategoryCache, base_id: int, report: Report) -> int:
    """Columns A-G: 日期, 金额, 币种, 内容, 分类, 人民币换算(workbook), 当日汇率(workbook)."""
    ws = wb[sheet_name]
    imported = 0
    for r in range(2, ws.max_row + 1):
        date_v = ws.cell(row=r, column=1).value
        amount_v = ws.cell(row=r, column=2).value
        ccy_v = ws.cell(row=r, column=3).value
        content_v = ws.cell(row=r, column=4).value
        cat_v = ws.cell(row=r, column=5).value
        wb_rmb = ws.cell(row=r, column=6).value
        wb_rate = ws.cell(row=r, column=7).value

        cells = [date_v, amount_v, ccy_v, content_v, cat_v, wb_rmb, wb_rate]
        if all(c in (None, "") for c in cells):
            continue  # truly blank row

        iso = _iso(date_v)
        amount_dec = money.parse_decimal(amount_v)
        ccy_row = ccy.get(ccy_v) if ccy_v is not None else None

        # --- Anomaly gate: must have a date, a parseable amount, and a currency.
        problems = []
        if iso is None:
            problems.append("missing/invalid date")
        if amount_dec is None:
            problems.append(f"unparseable amount: {amount_v!r}")
        if ccy_row is None:
            problems.append(f"unknown/missing currency: {ccy_v!r}")
        if problems:
            report.anomalies.append({
                "sheet": sheet_name, "row": r, "reasons": problems,
                "raw": {"date": str(date_v), "amount": str(amount_v),
                        "currency": str(ccy_v), "content": str(content_v)},
            })
            continue

        content = "" if content_v is None else str(content_v).strip()
        minor_unit = ccy_row["minor_unit"]
        amount_minor = money.to_minor(amount_dec, minor_unit)
        category_id = cats.get_or_create(cat_v, direction)
        cat_name = "" if cat_v is None else str(cat_v).strip()

        # --- Conversion via our own rate engine.
        rl = rates.lookup(conn, ccy_row["id"], iso, base_currency_id=base_id)
        rate_used = rate_date_used = base_micro = None
        stale_days = 0
        if rl is not None:
            rate_used = rl.rate_to_base
            rate_date_used = rl.rate_date
            stale_days = rl.stale_days
            base_micro = money.to_base_micro(amount_minor, minor_unit, rate_used)
        else:
            report.no_rate.append({"sheet": sheet_name, "row": r,
                                   "currency": ccy_row["code"], "date": iso})

        row_hash = _row_hash(direction, iso, amount_minor, ccy_row["code"],
                             content, cat_name, "")
        external_uid = make_external_uid(sheet_name, r, row_hash)

        conn.execute(
            """INSERT INTO txn (
                   external_uid, txn_date, direction, amount_minor, currency_id,
                   content, category_id, note, is_aa, payer_id, aa_settled,
                   rate_used, rate_date_used, base_amount_micro, rate_stale_days,
                   revision, last_xlsx_hash, source_sheet, source_row)
               VALUES (?, ?, ?, ?, ?, ?, ?, NULL, 0, NULL, 0,
                       ?, ?, ?, ?, 1, ?, ?, ?)""",
            (external_uid, iso, direction, amount_minor, ccy_row["id"],
             content, category_id, rate_used, rate_date_used, base_micro,
             stale_days, row_hash, sheet_name, r),
        )
        imported += 1

        # --- Post-import flags / reconciliation.
        if content and "AA" in content:
            report.aa_candidates.append({"sheet": sheet_name, "row": r,
                                         "content": content, "uid": external_uid})

        if rl is not None and rates.is_stale(stale_days):
            report.stale_rate.append({"sheet": sheet_name, "row": r, "date": iso,
                                      "currency": ccy_row["code"],
                                      "rate_date": rate_date_used,
                                      "stale_days": stale_days})

        # Reconcile against workbook's own 当日汇率 / 人民币换算 values.
        wb_rate_dec = money.parse_decimal(wb_rate)
        wb_rmb_dec = money.parse_decimal(wb_rmb)
        if rl is not None and wb_rmb_dec is None:
            # workbook never computed RMB (e.g. the income refund rows) -> we did.
            report.computed_rmb.append({"sheet": sheet_name, "row": r,
                                        "content": content,
                                        "computed_cny": str(money.base_micro_to_cny(base_micro))})
        elif rl is not None:
            our_rate = money.parse_decimal(rate_used)
            our_cny = money.base_micro_to_cny(base_micro)
            rate_diff = abs(float(our_rate) - float(wb_rate_dec)) if wb_rate_dec is not None else 0.0
            rmb_diff = abs(float(our_cny) - float(wb_rmb_dec))
            if rate_diff > config.RECONCILE_RATE_TOLERANCE or rmb_diff > config.RECONCILE_RMB_TOLERANCE:
                report.reconcile_diffs.append({
                    "sheet": sheet_name, "row": r, "content": content,
                    "our_rate": rate_used, "wb_rate": str(wb_rate_dec),
                    "our_cny": str(our_cny), "wb_cny": str(wb_rmb_dec),
                    "rate_diff": round(rate_diff, 6), "rmb_diff": round(rmb_diff, 4),
                })
    conn.commit()
    return imported


def import_workbook(conn: sqlite3.Connection,
                    workbook_path: Optional[Path] = None) -> Report:
    path = Path(workbook_path) if workbook_path else config.WORKBOOK_PATH
    # Not read_only: we use random cell access (ws.cell), unsupported in read_only mode.
    wb = openpyxl.load_workbook(str(path), data_only=True)
    report = Report()
    try:
        report.rates_loaded = load_rates(conn, wb)
        ccy = currency_id_map(conn)
        report.fiscal_months_loaded = load_fiscal_months(conn, wb, ccy)
        cats = _CategoryCache(conn)
        base_id = rates.get_base_currency_id(conn)
        report.expense_imported = _import_sheet(
            conn, wb, EXPENSE_SHEET, "expense", ccy, cats, base_id, report)
        report.income_imported = _import_sheet(
            conn, wb, INCOME_SHEET, "income", ccy, cats, base_id, report)
    finally:
        wb.close()
    return report
