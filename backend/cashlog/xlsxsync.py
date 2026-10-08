"""Two-way xlsx sync (P2).

  export_to_xlsx  —  "同步到表格": DB -> xlsx. Rewrites the 支出/入账 data sheets
                     from the DB (source of truth), writes a hidden _cashlog_uid
                     column, rebuilds 汇率/统计 so the workbook stays usable on
                     its own, backs up the workbook, saves atomically, then marks
                     the synced baseline on each txn. Calculator sheets are
                     preserved untouched.

  diff_from_xlsx  —  "从表格同步" preview: xlsx -> proposed changes. Matches rows
                     by _cashlog_uid (so row reorder is a no-op) and by content
                     hash (so unchanged rows are skipped). Classifies each row as
                     new / modified / conflict / unknown_uid / anomaly. A DB row
                     missing from the sheet is a deletion only if it was synced
                     before (last_synced_at set); an app-created row not yet
                     exported goes to db_only, never deleted. Applies nothing.

  apply_diff      —  Applies new + modified (deletions/conflicts are opt-in).

Identity & conflict model (per txn):
  last_xlsx_hash = content hash at the last successful sync (baseline).
  db_changed     = current DB hash   != baseline
  xlsx_changed   = incoming xlsx hash != baseline
  conflict       = both changed AND they disagree -> never auto-overwritten.
"""
from __future__ import annotations

import os
import shutil
import sqlite3
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl
from openpyxl.styles import Alignment

from . import config, importer, jsonsync, money, rates, stats

# direction -> sheet name + the workbook's column-D header
DATA_SHEETS = {
    "expense": ("支出", "消费内容"),
    "income": ("入账", "入账来源"),
}
UID_COL = 8  # column H
UID_HEADER = "_cashlog_uid"
# Excel built-in short date. Excel renders it with the user's regional settings
# (for this workbook that appears as e.g. 2026/6/24), unlike a literal ISO mask.
DATE_FMT = "mm-dd-yy"
RATE_SHEET = "汇率"
STATS_SHEET = "统计"


class WorkbookLockedError(RuntimeError):
    """Raised when the target xlsx cannot be replaced (typically open in Excel)."""

    def __init__(self, path):
        super().__init__(
            f"无法写入 {path}：文件可能正被 Excel 打开。请关闭后重试。"
        )
        self.path = path


def _now() -> str:
    return datetime.now().astimezone().strftime("%Y-%m-%dT%H:%M:%S")


def _header(content_header: str) -> list:
    return ["日期", "金额", "币种", content_header, "分类",
            "人民币换算", "当日汇率", UID_HEADER]


def _db_hash(direction: str, txn_date: str, amount_minor: int, code: str,
             content: str, cat_name: Optional[str], note: Optional[str] = None) -> str:
    """Row hash over exactly what a sheet row can carry. The note has no column,
    so it is left out (a txn with a note used to diff as "modified" right after
    a clean export); content / category are stripped the same way the diff
    reads them back. ``note`` is accepted for call-site compatibility only."""
    return importer._row_hash(direction, txn_date, amount_minor, code,
                              (content or "").strip(), (cat_name or "").strip(), "")


def _fetch_txns(conn: sqlite3.Connection, direction: str) -> list:
    """Rows to sync. ``eff_amount_minor`` (txn_eff view) is what the workbook
    holds: my own share for a settled AA expense, the full amount otherwise.
    The row hash is always taken over that workbook-facing amount, so export and
    diff agree; ``amount_minor`` (the stored total) is never written to the sheet
    for those rows and is restored when an edit comes back from it."""
    return conn.execute(
        """SELECT t.id, t.external_uid, t.txn_date, t.amount_minor, t.content,
                  t.eff_amount_minor, t.eff_is_share,
                  t.note, t.rate_used, t.last_xlsx_hash, t.last_synced_at, t.revision,
                  cur.code AS code, cur.name AS ccy_name, cur.minor_unit,
                  c.name AS cat_name
           FROM   txn_eff t
           JOIN   currency cur ON cur.id = t.currency_id
           LEFT   JOIN category c ON c.id = t.category_id
           WHERE  t.direction = ?
           ORDER  BY t.txn_date, t.id""",
        (direction,),
    ).fetchall()


def _safe_unlink(path: Path) -> None:
    try:
        if path.exists():
            path.unlink()
    except OSError:
        pass


def _clear_trailing(ws, first_row: int, last_row: int, ncols: int,
                    templates: list | None = None, *,
                    date_cols: set[int] | None = None) -> None:
    date_cols = date_cols or set()
    for rr in range(first_row, last_row + 1):
        for col in range(1, ncols + 1):
            _write_cell(ws, rr, col, None, templates, date_cell=(col in date_cols))


def _centered(alignment) -> Alignment:
    """Preserve existing alignment details, filling missing axes with center."""
    return Alignment(
        horizontal=alignment.horizontal or "center",
        vertical=alignment.vertical or "center",
        text_rotation=alignment.textRotation,
        wrap_text=alignment.wrap_text,
        shrink_to_fit=alignment.shrink_to_fit,
        indent=alignment.indent,
        relativeIndent=alignment.relativeIndent,
        justifyLastLine=alignment.justifyLastLine,
        readingOrder=alignment.readingOrder,
    )


def _column_style_templates(ws, ncols: int, *,
                            preferred_rows: tuple[int, ...] = (2, 1)) -> list:
    """Snapshot per-column styles before rewriting rows.

    ``openpyxl`` creates blank cells with default left-ish display. When export
    adds rows beyond the existing workbook range, copying a nearby column style
    keeps the workbook looking like it did before the app touched it.
    """
    templates = []
    for col in range(1, ncols + 1):
        picked = None
        for row in preferred_rows:
            cell = ws.cell(row=row, column=col)
            if cell.has_style:
                picked = copy(cell._style)
                break
        templates.append(picked)
    return templates


def _write_cell(ws, row: int, col: int, value: Any, templates: list | None = None, *,
                date_cell: bool = False):
    cell = ws.cell(row=row, column=col)
    cell.value = value
    if templates and col <= len(templates) and templates[col - 1] is not None:
        cell._style = copy(templates[col - 1])
    cell.alignment = _centered(cell.alignment)
    if date_cell:
        cell.number_format = DATE_FMT
    return cell


def _rebuild_rates_sheet(conn: sqlite3.Connection, wb) -> None:
    """汇率: 日期 / 币种 / 当日汇率, rebuilt from exchange_rate (values)."""
    ws = wb[RATE_SHEET] if RATE_SHEET in wb.sheetnames else wb.create_sheet(RATE_SHEET)
    original_max = ws.max_row
    templates = _column_style_templates(ws, 3)
    for col, h in enumerate(["日期", "币种", "当日汇率"], start=1):
        _write_cell(ws, 1, col, h, templates, date_cell=(col == 1))
    r = 2
    for row in conn.execute(
        """SELECT er.rate_date, cur.name AS ccy, er.rate_to_base
           FROM exchange_rate er JOIN currency cur ON cur.id = er.currency_id
           ORDER BY er.rate_date, cur.code"""):
        _write_cell(ws, r, 1, date.fromisoformat(row["rate_date"]), templates, date_cell=True)
        _write_cell(ws, r, 2, row["ccy"], templates)
        _write_cell(ws, r, 3, float(row["rate_to_base"]), templates)
        r += 1
    _clear_trailing(ws, r, original_max, 3, templates, date_cols={1})


def _rebuild_stats_sheet(conn: sqlite3.Connection, wb, base_id: int,
                         usd_id: Optional[int]) -> None:
    """统计: rebuilt from fiscal_month. Derived columns use the same formulas as
    the original workbook (so the file stays standalone-usable); 当日汇率 and the
    manual 信用卡金额 are written as values. The 学费/特殊 exclusion list is taken
    live from the DB so the SUMIFS stays correct if categories change."""
    excl = [row["name"] for row in conn.execute(
        "SELECT name FROM category WHERE kind='expense' AND exclude_stats=1 ORDER BY name")]
    excl_terms = "".join(f', 支出!$E:$E, "<>{name}"' for name in excl)

    ws = wb[STATS_SHEET] if STATS_SHEET in wb.sheetnames else wb.create_sheet(STATS_SHEET)
    original_max = ws.max_row
    headers = ["财政月", "总支出金额", "币种", "美元换算", "信用卡金额",
               "币种", "人民币换算", "当日汇率", "占总支出"]
    templates = _column_style_templates(ws, 9)
    for col, h in enumerate(headers, start=1):
        _write_cell(ws, 1, col, h, templates, date_cell=(col == 1))

    # Same months as the app's 统计 page (auto-filled up to the current one).
    ccy_names = {row["code"]: row["name"] for row in conn.execute("SELECT code, name FROM currency")}
    r = 2
    for fm in stats.fiscal_months(conn):
        pe = fm["period_end"]
        usd_rate = fm["usd_rate"]
        cc_code = fm["credit_card_currency"] or "USD"
        cc_major = fm["credit_card"]
        total_override = fm["total_cny_override"]
        _write_cell(ws, r, 1, date.fromisoformat(pe), templates, date_cell=True)
        if total_override is None:
            _write_cell(ws, r, 2,
                        f'=SUMIFS(支出!$F:$F, 支出!$A:$A,">="&EDATE($A{r},-1), '
                        f'支出!$A:$A,"<"&$A{r}{excl_terms})', templates)
        else:
            _write_cell(ws, r, 2, float(total_override), templates)
        _write_cell(ws, r, 3, "人民币", templates)
        _write_cell(ws, r, 4, f"=B{r}/H{r}", templates)
        _write_cell(ws, r, 5, cc_major, templates)
        _write_cell(ws, r, 6, ccy_names.get(cc_code, cc_code), templates)
        if cc_code == "USD":
            _write_cell(ws, r, 7, f"=E{r}*H{r}", templates)
        else:  # card in another currency: its own rate, not the USD one
            cc_cny = fm["credit_card_cny"]
            _write_cell(ws, r, 7, cc_cny, templates)
        _write_cell(ws, r, 8, usd_rate, templates)
        _write_cell(ws, r, 9, f"=G{r}/B{r}", templates)
        r += 1
    _clear_trailing(ws, r, original_max, 9, templates, date_cols={1})


# ---------------------------------------------------------------------------
# Export: DB -> xlsx
# ---------------------------------------------------------------------------
class PendingSheetChangesError(Exception):
    """The workbook holds edits not yet imported; exporting would discard them."""

    def __init__(self, summary: dict):
        self.summary = summary
        parts = [f"{label} {summary[k]}" for k, label in
                 (("new", "新增"), ("modified", "修改"), ("deleted", "删除"),
                  ("conflict", "冲突"), ("unknown_uid", "未知UID"), ("anomalies", "异常"))
                 if summary.get(k)]
        super().__init__("表格里有尚未导入的改动（" + "，".join(parts) +
                         "）。直接同步到表格会覆盖它们；请先「从表格同步」，或确认后强制覆盖。")


def pending_sheet_changes(diff: "SyncDiff") -> int:
    s = diff.summary()
    return sum(s[k] for k in ("new", "modified", "deleted", "conflict", "unknown_uid", "anomalies"))


def export_to_xlsx(conn: sqlite3.Connection, target: Optional[Path] = None, *,
                   template: Optional[Path] = None, backup: bool = True,
                   backup_dir: Optional[Path] = None, write_json: bool = True,
                   force: bool = False) -> dict:
    """DB -> workbook. The data sheets are rewritten from the DB, so unless
    ``force`` is set this first diffs the existing target and raises
    PendingSheetChangesError when it holds edits that were never imported
    (they would otherwise be silently wiped; a backup is still taken)."""
    target = Path(target) if target else config.WORKBOOK_PATH
    load_from = Path(template) if template else (target if target.exists() else None)
    if not force and template is None and target.exists():
        diff = diff_from_xlsx(conn, target)
        if pending_sheet_changes(diff):
            raise PendingSheetChangesError(diff.summary())

    if load_from and load_from.exists():
        wb = openpyxl.load_workbook(str(load_from))  # keep formulas/other sheets
    else:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    pending: list = []  # (txn_id, db_hash) to mark as synced AFTER a clean write
    written = {"expense": 0, "income": 0}

    for direction, (sheet_name, content_header) in DATA_SHEETS.items():
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
        original_max = ws.max_row
        templates = _column_style_templates(ws, UID_COL)

        for col, value in enumerate(_header(content_header), start=1):
            _write_cell(ws, 1, col, value, templates, date_cell=(col == 1))

        r = 2
        for t in _fetch_txns(conn, direction):
            major = money.from_minor(t["eff_amount_minor"], t["minor_unit"])
            _write_cell(ws, r, 1, date.fromisoformat(t["txn_date"]), templates, date_cell=True)
            _write_cell(ws, r, 2, float(major), templates)
            _write_cell(ws, r, 3, t["ccy_name"], templates)
            _write_cell(ws, r, 4, t["content"], templates)
            _write_cell(ws, r, 5, t["cat_name"] or "", templates)
            _write_cell(ws, r, 6, f"=B{r}*G{r}", templates)
            _write_cell(ws, r, 7, float(t["rate_used"]) if t["rate_used"] is not None else None, templates)
            _write_cell(ws, r, UID_COL, t["external_uid"], templates)
            pending.append((t["id"], _db_hash(direction, t["txn_date"], t["eff_amount_minor"],
                                              t["code"], t["content"], t["cat_name"], t["note"])))
            r += 1
            written[direction] += 1

        # NB: ws.cell(..., value=None) is a no-op in openpyxl; _clear_trailing
        # assigns .value on the cell object to actually clear it.
        _clear_trailing(ws, r, original_max, UID_COL, templates, date_cols={1})

        ws.column_dimensions[openpyxl.utils.get_column_letter(UID_COL)].hidden = True

    base_id = rates.get_base_currency_id(conn)
    usd = conn.execute("SELECT id FROM currency WHERE code='USD'").fetchone()
    _rebuild_rates_sheet(conn, wb)
    _rebuild_stats_sheet(conn, wb, base_id, usd["id"] if usd else None)

    # --- Backup, then atomic save. DB is only marked synced on success. ---
    if backup and target.exists():
        bdir = Path(backup_dir) if backup_dir else config.BACKUP_DIR
        bdir.mkdir(parents=True, exist_ok=True)
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(target, bdir / f"{target.stem}_{ts}{target.suffix}")

    tmp = target.with_name(target.name + ".tmp")
    try:
        wb.save(str(tmp))
        os.replace(str(tmp), str(target))
    except PermissionError as exc:
        _safe_unlink(tmp)
        raise WorkbookLockedError(target) from exc
    except Exception:
        _safe_unlink(tmp)
        raise

    now = _now()
    for txn_id, h in pending:
        conn.execute(
            "UPDATE txn SET last_xlsx_hash = ?, last_synced_at = ? WHERE id = ?",
            (h, now, txn_id),
        )
    conn.commit()
    if write_json:
        jsonsync.export_json(conn)

    return {"target": str(target), "written": written, "synced_at": now}


# ---------------------------------------------------------------------------
# Diff: xlsx -> proposed changes (no writes)
# ---------------------------------------------------------------------------
@dataclass
class SyncDiff:
    new: list = field(default_factory=list)
    modified: list = field(default_factory=list)
    deleted: list = field(default_factory=list)
    conflict: list = field(default_factory=list)
    unknown_uid: list = field(default_factory=list)
    anomalies: list = field(default_factory=list)
    db_only: list = field(default_factory=list)   # in DB, never synced to xlsx yet
    unchanged: int = 0
    # Rows already in the DB whose uid is missing from the sheet (written back on apply).
    uid_backfill: list = field(default_factory=list)

    def summary(self) -> dict:
        return {
            "new": len(self.new), "modified": len(self.modified),
            "deleted": len(self.deleted), "conflict": len(self.conflict),
            "unknown_uid": len(self.unknown_uid), "anomalies": len(self.anomalies),
            "db_only": len(self.db_only), "unchanged": self.unchanged,
        }


def _iso(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return None


def diff_from_xlsx(conn: sqlite3.Connection, source: Optional[Path] = None) -> SyncDiff:
    source = Path(source) if source else config.WORKBOOK_PATH
    wb = openpyxl.load_workbook(str(source), data_only=True)
    diff = SyncDiff()

    ccy = importer.currency_id_map(conn)
    cat_map = {(row["name"], row["kind"]): row["id"]
               for row in conn.execute("SELECT id, name, kind FROM category")}

    # DB snapshot keyed by uid, with current hash + synced baseline.
    db_by_uid: dict = {}
    in_scope_uids = {"expense": set(), "income": set()}
    for direction in DATA_SHEETS:
        for t in _fetch_txns(conn, direction):
            cur_hash = _db_hash(direction, t["txn_date"], t["eff_amount_minor"], t["code"],
                                t["content"], t["cat_name"])
            db_by_uid[t["external_uid"]] = {
                "id": t["id"], "direction": direction, "current_hash": cur_hash,
                "is_share": bool(t["eff_is_share"]), "code": t["code"],
                "amount_minor": t["amount_minor"], "eff_amount_minor": t["eff_amount_minor"],
                "baseline": t["last_xlsx_hash"], "content": t["content"],
                "synced": t["last_synced_at"] is not None,
            }
            in_scope_uids[direction].add(t["external_uid"])

    seen: set = set()
    try:
        for direction, (sheet_name, _) in DATA_SHEETS.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            for r in range(2, ws.max_row + 1):
                cells = [ws.cell(row=r, column=c).value for c in range(1, 6)]
                uid = ws.cell(row=r, column=UID_COL).value
                uid = str(uid).strip() if uid not in (None, "") else None
                if all(c in (None, "") for c in cells) and not uid:
                    continue

                # A uid seen earlier in the sheet means this row was copy-pasted
                # together with the hidden uid column: treat the copy as a new
                # row instead of letting it overwrite the original txn.
                if uid and uid in seen:
                    uid = None
                if uid:
                    # Mark it seen before any anomaly check, so a row with a bad
                    # edit is reported as an anomaly, never also as a deletion.
                    seen.add(uid)

                def anomaly(*reasons):
                    diff.anomalies.append({"sheet": sheet_name, "row": r,
                                           "reasons": list(reasons), "uid": uid})

                iso = _iso(cells[0])
                amount_dec = money.parse_decimal(cells[1])
                ccy_row = ccy.get(cells[2]) if cells[2] is not None else None
                problems = []
                if iso is None:
                    problems.append("missing/invalid date")
                if amount_dec is None:
                    problems.append(f"unparseable amount: {cells[1]!r}")
                if ccy_row is None:
                    problems.append(f"unknown/missing currency: {cells[2]!r}")
                if problems:
                    anomaly(*problems)
                    continue
                try:
                    amount_minor = money.to_minor(amount_dec, ccy_row["minor_unit"])
                except ValueError as exc:
                    anomaly(str(exc))
                    continue

                content = "" if cells[3] is None else str(cells[3]).strip()
                cat_name = None if cells[4] in (None, "") else str(cells[4]).strip()
                cat_id = cat_map.get((cat_name, direction)) if cat_name else None
                if cat_name and cat_id is None:
                    # Unknown name (or one of the other kind): applying would
                    # silently drop it to 未分类 and the next export would erase it.
                    anomaly(f"未知分类「{cat_name}」：请先在应用中添加该分类，或改成已有分类")
                    continue
                xlsx_hash = _db_hash(direction, iso, amount_minor, ccy_row["code"],
                                     content, cat_name)
                payload = {"sheet": sheet_name, "row": r, "uid": uid, "direction": direction,
                           "txn_date": iso, "amount_minor": amount_minor,
                           "currency_id": ccy_row["id"], "content": content,
                           "category_id": cat_id, "xlsx_hash": xlsx_hash}

                if not uid:
                    # Imported earlier (e.g. by migrate or a previous apply) but
                    # the uid never made it into the sheet: the deterministic uid
                    # recognises it instead of inserting a duplicate.
                    known = importer.make_external_uid(sheet_name, r, xlsx_hash)
                    rec = db_by_uid.get(known)
                    if rec is not None and known not in seen and rec["direction"] == direction:
                        seen.add(known)
                        diff.unchanged += 1
                        diff.uid_backfill.append({"sheet": sheet_name, "row": r, "uid": known})
                        continue
                    diff.new.append(payload)
                    continue

                rec = db_by_uid.get(uid)
                if rec is None:
                    diff.unknown_uid.append({"sheet": sheet_name, "row": r, "uid": uid})
                    continue

                if rec["direction"] != direction:
                    anomaly("该行被移到了另一张表（支出/入账）：请在应用中修改收支方向")
                    continue

                if xlsx_hash == rec["current_hash"]:
                    # Sheet already matches the DB (whatever the baseline says).
                    diff.unchanged += 1
                    continue

                if rec["is_share"] and xlsx_hash != rec["baseline"]:
                    # Structured AA: the sheet shows my share, but the DB must keep
                    # the full amount (the split is derived from it). An amount /
                    # currency edit made in Excel can't be mapped back to a total,
                    # so it is reported instead of silently overwriting it; any
                    # other edit applies with the stored total restored. (A row
                    # untouched in Excel — e.g. still showing a pre-AA total — is
                    # just "unchanged"; the next export rewrites it.)
                    if (amount_minor != rec["eff_amount_minor"]
                            or ccy_row["code"] != rec["code"]):
                        anomaly("AA 交易的金额/币种请在应用内修改（表格中显示的是我的部分）")
                        continue
                    payload["amount_minor"] = rec["amount_minor"]

                xlsx_changed = xlsx_hash != rec["baseline"]
                db_changed = rec["current_hash"] != rec["baseline"]
                if not xlsx_changed:
                    diff.unchanged += 1
                elif db_changed:
                    diff.conflict.append({**payload, "txn_id": rec["id"],
                                          "db_content": rec["content"]})
                else:
                    diff.modified.append({**payload, "txn_id": rec["id"]})

        for direction in DATA_SHEETS:
            for uid in in_scope_uids[direction] - seen:
                rec = db_by_uid[uid]
                entry = {"uid": uid, "txn_id": rec["id"],
                         "direction": direction, "content": rec["content"]}
                # Only rows that were previously written to the workbook can be a
                # deletion. A row created in the app but not yet "同步到表格" is
                # simply unsynced — never a delete candidate.
                if rec["synced"]:
                    diff.deleted.append(entry)
                else:
                    diff.db_only.append(entry)
    finally:
        wb.close()
    return diff


# ---------------------------------------------------------------------------
# Apply
# ---------------------------------------------------------------------------
def _convert(conn, currency_id: int, txn_date: str, amount_minor: int, base_id: int):
    cur = conn.execute("SELECT minor_unit FROM currency WHERE id = ?", (currency_id,)).fetchone()
    rl = rates.lookup(conn, currency_id, txn_date, base_currency_id=base_id)
    if rl is None:
        return None, None, 0, None
    base_micro = money.to_base_micro(amount_minor, cur["minor_unit"], rl.rate_to_base)
    return rl.rate_to_base, rl.rate_date, rl.stale_days, base_micro


def _write_uids(path: Path, writes: list) -> None:
    """Set the hidden uid cell of the given (sheet, row) pairs; atomic save."""
    wb = openpyxl.load_workbook(str(path))
    for sheet, row, uid in writes:
        if sheet in wb.sheetnames:
            wb[sheet].cell(row=row, column=UID_COL).value = uid
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(str(tmp))
        os.replace(str(tmp), str(path))
    except PermissionError as exc:
        _safe_unlink(tmp)
        raise WorkbookLockedError(path) from exc
    except Exception:
        _safe_unlink(tmp)
        raise


def apply_diff(conn: sqlite3.Connection, source: Optional[Path], diff: SyncDiff, *,
               apply_deletions: bool = False, apply_conflicts: bool = False,
               write_json: bool = True) -> dict:
    base_id = rates.get_base_currency_id(conn)
    now = _now()
    applied = {"new": 0, "modified": 0, "deleted": 0, "conflict": 0}

    def upsert_values(p):
        rate_used, rate_date, stale, base_micro = _convert(
            conn, p["currency_id"], p["txn_date"], p["amount_minor"], base_id)
        return rate_used, rate_date, stale, base_micro

    for p in diff.new:
        rate_used, rate_date, stale, base_micro = upsert_values(p)
        uid = importer.make_external_uid(p["sheet"], p["row"], p["xlsx_hash"])
        conn.execute(
            """INSERT INTO txn (external_uid, txn_date, direction, amount_minor,
                   currency_id, content, category_id, is_aa, aa_settled,
                   rate_used, rate_date_used, base_amount_micro, rate_stale_days,
                   revision, last_xlsx_hash, last_synced_at, source_sheet, source_row)
               VALUES (?, ?, ?, ?, ?, ?, ?, 0, 0, ?, ?, ?, ?, 1, ?, ?, ?, ?)""",
            (uid, p["txn_date"], p["direction"], p["amount_minor"], p["currency_id"],
             p["content"], p["category_id"], rate_used, rate_date, base_micro, stale,
             p["xlsx_hash"], now, p["sheet"], p["row"]))
        applied["new"] += 1

    targets = list(diff.modified) + (list(diff.conflict) if apply_conflicts else [])
    uid_writes = [(p["sheet"], p["row"], importer.make_external_uid(p["sheet"], p["row"], p["xlsx_hash"]))
                  for p in diff.new]
    uid_writes += [(b["sheet"], b["row"], b["uid"]) for b in diff.uid_backfill]
    for p in targets:
        rate_used, rate_date, stale, base_micro = upsert_values(p)
        conn.execute(
            """UPDATE txn SET txn_date=?, amount_minor=?, currency_id=?, content=?,
                   category_id=?, rate_used=?, rate_date_used=?, base_amount_micro=?,
                   rate_stale_days=?, revision=revision+1, last_xlsx_hash=?, last_synced_at=?
               WHERE id=?""",
            (p["txn_date"], p["amount_minor"], p["currency_id"], p["content"],
             p["category_id"], rate_used, rate_date, base_micro, stale,
             p["xlsx_hash"], now, p["txn_id"]))
        applied["conflict" if p in diff.conflict else "modified"] += 1

    if apply_deletions:
        for d in diff.deleted:
            conn.execute("DELETE FROM txn WHERE id = ?", (d["txn_id"],))
            applied["deleted"] += 1

    # Imported rows get their uid written into the hidden column right away, so a
    # second import before the next export recognises them instead of inserting
    # duplicates (or failing on the uid UNIQUE constraint). Sheet first, then the
    # DB commit: if the workbook can't be saved (open in Excel) nothing is applied.
    if uid_writes:
        try:
            _write_uids(Path(source) if source else config.WORKBOOK_PATH, uid_writes)
        except Exception:
            conn.rollback()
            raise
    conn.commit()
    if write_json:
        jsonsync.export_json(conn)
    return applied
