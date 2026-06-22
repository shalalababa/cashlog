"""P2 xlsx round-trip sync — acceptance scenarios. Fully isolated in a temp dir;
the real 记账单.xlsx is never touched. Runnable with pytest or directly."""
from __future__ import annotations

import sys
import tempfile
import uuid
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from cashlog import db, seed, xlsxsync  # noqa: E402


def _build_env():
    d = Path(tempfile.mkdtemp(prefix="cashlog_p2_"))
    dbp = d / "t.db"
    db.init_db(dbp, reset=True)
    conn = db.connect(dbp)
    seed.seed(conn)
    cny = conn.execute("SELECT id FROM currency WHERE code='CNY'").fetchone()["id"]
    ecat = conn.execute("SELECT id FROM category WHERE name='饮食' AND kind='expense'").fetchone()["id"]
    icat = conn.execute("SELECT id FROM category WHERE name='利息' AND kind='income'").fetchone()["id"]

    def add(direction, date, minor, content, catid):
        conn.execute(
            """INSERT INTO txn (external_uid, txn_date, direction, amount_minor,
                   currency_id, content, category_id, rate_used, base_amount_micro, revision)
               VALUES (?, ?, ?, ?, ?, ?, ?, '1', ?, 1)""",
            (uuid.uuid4().hex, date, direction, minor, cny, content, catid, minor * 10000))

    add("expense", "2024-01-10", 1000, "咖啡", ecat)
    add("expense", "2024-01-11", 2550, "午饭", ecat)
    add("expense", "2024-01-12", 999, "地铁", ecat)
    add("income", "2024-01-15", 500, "利息", icat)
    conn.commit()
    xlsx = d / "book.xlsx"
    return d, conn, xlsx


def _export(conn, xlsx):
    return xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)


def test_round_trip_is_clean():
    _, conn, xlsx = _build_env()
    _export(conn, xlsx)
    s = xlsxsync.diff_from_xlsx(conn, xlsx).summary()
    assert s["new"] == 0 and s["modified"] == 0 and s["conflict"] == 0
    assert s["deleted"] == 0 and s["anomalies"] == 0
    assert s["unchanged"] == 4
    # The uid column is written and hidden.
    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb["支出"]
    assert ws.cell(row=1, column=8).value == xlsxsync.UID_HEADER
    assert ws.column_dimensions["H"].hidden is True


def test_row_reorder_is_not_a_change():
    _, conn, xlsx = _build_env()
    _export(conn, xlsx)
    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb["支出"]
    row2 = [ws.cell(row=2, column=c).value for c in range(1, 9)]
    row3 = [ws.cell(row=3, column=c).value for c in range(1, 9)]
    for c in range(1, 9):
        ws.cell(row=2, column=c, value=row3[c - 1])
        ws.cell(row=3, column=c, value=row2[c - 1])
    wb.save(str(xlsx))

    s = xlsxsync.diff_from_xlsx(conn, xlsx).summary()
    assert s["new"] == 0, "uid matching must not see reorder as a new row"
    assert s["modified"] == 0 and s["conflict"] == 0
    assert s["unchanged"] == 4


def test_single_modify_shows_one_change():
    _, conn, xlsx = _build_env()
    _export(conn, xlsx)
    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb["支出"]
    ws.cell(row=2, column=2, value=12.00)  # was 10.00
    wb.save(str(xlsx))

    diff = xlsxsync.diff_from_xlsx(conn, xlsx)
    assert diff.summary()["modified"] == 1
    assert diff.summary()["unchanged"] == 3
    xlsxsync.apply_diff(conn, xlsx, diff, write_json=False)
    amt = conn.execute(
        "SELECT amount_minor FROM txn WHERE id=?", (diff.modified[0]["txn_id"],)
    ).fetchone()["amount_minor"]
    assert amt == 1200


def test_both_sides_modified_is_a_conflict():
    _, conn, xlsx = _build_env()
    _export(conn, xlsx)
    # DB-side edit (no re-sync, so baseline stays put -> db_changed).
    conn.execute("UPDATE txn SET content='咖啡改DB' WHERE content='咖啡'")
    conn.commit()
    # xlsx-side edit on the same row.
    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb["支出"]
    ws.cell(row=2, column=2, value=99.00)
    wb.save(str(xlsx))

    diff = xlsxsync.diff_from_xlsx(conn, xlsx)
    assert diff.summary()["conflict"] == 1
    assert diff.summary()["modified"] == 0
    # Default apply must NOT overwrite a conflict.
    xlsxsync.apply_diff(conn, xlsx, diff, write_json=False)
    row = conn.execute(
        "SELECT content, amount_minor FROM txn WHERE id=?", (diff.conflict[0]["txn_id"],)
    ).fetchone()
    assert row["content"] == "咖啡改DB" and row["amount_minor"] == 1000


def test_export_clears_stale_template_rows():
    # A template with junk rows beyond the DB's data must be wiped on export,
    # otherwise they resurface as phantom new/anomaly rows on the next diff.
    _, conn, xlsx = _build_env()
    _export(conn, xlsx)
    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb["支出"]
    ws.cell(row=8, column=1, value=None)
    ws.cell(row=8, column=4).value = "幽灵行"      # junk content, no uid
    ws.cell(row=9, column=4).value = "另一条幽灵"
    wb.save(str(xlsx))

    _export(conn, xlsx)  # re-export should clear rows 5..9
    s = xlsxsync.diff_from_xlsx(conn, xlsx).summary()
    assert s["new"] == 0 and s["anomalies"] == 0
    assert s["unchanged"] == 4


def test_write_failure_leaves_db_unmarked():
    d, conn, xlsx = _build_env()
    _export(conn, xlsx)
    before = conn.execute(
        "SELECT last_synced_at FROM txn ORDER BY id LIMIT 1").fetchone()["last_synced_at"]
    assert before is not None
    bad_target = d / "missing_dir" / "out.xlsx"  # parent does not exist -> save fails
    try:
        xlsxsync.export_to_xlsx(conn, target=bad_target, template=xlsx,
                                backup=False, write_json=False)
        raised = False
    except Exception:
        raised = True
    assert raised, "export to an unwritable target should raise"
    after = conn.execute(
        "SELECT last_synced_at FROM txn ORDER BY id LIMIT 1").fetchone()["last_synced_at"]
    assert after == before, "failed write must not advance the synced baseline"


def test_locked_target_raises_friendly_error():
    d, conn, xlsx = _build_env()
    orig_replace = xlsxsync.os.replace

    def boom(_a, _b):
        raise PermissionError(13, "file is open in another process")

    xlsxsync.os.replace = boom
    try:
        raised = False
        try:
            xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)
        except xlsxsync.WorkbookLockedError as exc:
            raised = True
            assert "Excel" in str(exc)
        assert raised, "a locked target must raise WorkbookLockedError"
        assert not (xlsx.with_name(xlsx.name + ".tmp")).exists(), "tmp must be cleaned up"
        n = conn.execute(
            "SELECT COUNT(*) FROM txn WHERE last_synced_at IS NOT NULL").fetchone()[0]
        assert n == 0, "failed write must not advance the synced baseline"
    finally:
        xlsxsync.os.replace = orig_replace


def test_rebuild_rates_and_stats_sheets():
    d = Path(tempfile.mkdtemp(prefix="cashlog_stats_"))
    dbp = d / "t.db"
    db.init_db(dbp, reset=True)
    conn = db.connect(dbp)
    seed.seed(conn)
    usd = conn.execute("SELECT id FROM currency WHERE code='USD'").fetchone()["id"]
    ecat = conn.execute("SELECT id FROM category WHERE name='饮食' AND kind='expense'").fetchone()["id"]
    for dt, rate in (("2024-01-03", "7.10"), ("2024-01-05", "7.20")):
        conn.execute("INSERT INTO exchange_rate (currency_id, rate_date, rate_to_base) VALUES (?,?,?)",
                     (usd, dt, rate))
    # an expense inside the fiscal window [2023-12-05, 2024-01-05)
    conn.execute(
        """INSERT INTO txn (external_uid, txn_date, direction, amount_minor, currency_id,
               content, category_id, rate_used, base_amount_micro, revision)
           VALUES (?, '2023-12-20', 'expense', 1000, ?, '测试', ?, '7.10', 71000000, 1)""",
        (uuid.uuid4().hex, usd, ecat))
    conn.execute(
        """INSERT INTO fiscal_month (period_end, credit_card_amount_minor, credit_card_ccy_id)
           VALUES ('2024-01-05', 10000, ?)""", (usd,))
    conn.commit()

    xlsx = d / "book.xlsx"
    xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)
    wb = openpyxl.load_workbook(str(xlsx))

    rates_ws = wb["汇率"]
    assert [rates_ws.cell(row=1, column=c).value for c in range(1, 4)] == ["日期", "币种", "当日汇率"]
    assert rates_ws.cell(row=2, column=3).value == 7.10
    assert rates_ws.cell(row=3, column=3).value == 7.20

    st = wb["统计"]
    assert st.cell(row=2, column=8).value == 7.20         # 当日汇率 on the 5th
    assert st.cell(row=2, column=5).value == 100.0        # 信用卡金额 (USD)
    assert str(st.cell(row=2, column=2).value).startswith("=SUMIFS(")
    assert st.cell(row=2, column=4).value == "=B2/H2"
    assert st.cell(row=2, column=9).value == "=G2/B2"
    assert "学费" in st.cell(row=2, column=2).value and "特殊" in st.cell(row=2, column=2).value


def _run():
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_")]
    for fn in fns:
        fn()
        print(f"  ok: {fn.__name__}")
    print(f"{len(fns)} tests passed")


if __name__ == "__main__":
    _run()
