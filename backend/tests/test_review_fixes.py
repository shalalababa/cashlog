"""Regression tests for the repo-wide bug review (2026-10-08).

Each test pins one reproduced bug: atomic txn + AA writes, AA re-split on edit,
non-finite / huge numbers, receipt validation, fiscal USD rate & override, and
the xlsx sync round-trip issues. Isolated temp DB / workbook only.
"""
from __future__ import annotations

import sys
import tempfile
import uuid
from datetime import date
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from cashlog import aa_service, budgets, dashboard, db, money, rates, seed, service, stats, xlsxsync  # noqa: E402


def _env():
    d = Path(tempfile.mkdtemp(prefix="cashlog_fix_"))
    db.init_db(d / "t.db", reset=True)
    conn = db.connect(d / "t.db")
    seed.seed(conn)
    me = conn.execute("SELECT id FROM person WHERE is_self = 1").fetchone()["id"]
    bob = aa_service.create_person(conn, "Bob", write_json=False)["id"]
    return d, conn, me, bob


def _cat(conn, name, kind="expense"):
    return conn.execute("SELECT id FROM category WHERE name=? AND kind=?", (name, kind)).fetchone()["id"]


def _txn(conn, amount="100", **kw):
    data = {"date": "2024-03-10", "direction": "expense", "amount": amount,
            "currency": "CNY", "content": "x", **kw}
    return service.create_txn(conn, data, write_json=False)


def _count(conn):
    return conn.execute("SELECT COUNT(*) FROM txn").fetchone()[0]


# --------------------------------------------------------------- txn + AA atomicity
def test_create_with_aa_is_atomic():
    _, conn, me, bob = _env()
    with pytest.raises(ValueError):           # negative amount can't be AA
        _txn(conn, "-30", aa={"payer_id": me, "participants": [{"person_id": me}, {"person_id": bob}]})
    assert _count(conn) == 0                  # nothing half-written

    t = _txn(conn, "100", aa={"payer_id": me, "participants": [{"person_id": me}, {"person_id": bob}]})
    assert t["is_aa"] and t["my_amount"] == "50.00"


def test_loose_is_aa_fields_are_ignored():
    _, conn, me, _ = _env()
    t = _txn(conn, "30", is_aa=True, payer_id=me, aa_settled=True)
    assert not t["is_aa"] and t["my_amount"] is None
    assert stats.expense_total_micro(conn, "2024-03-05", "2024-04-05") == 30 * 1_000_000


def test_update_amount_resplits_equal_aa():
    _, conn, me, bob = _env()
    t = _txn(conn, "100", aa={"payer_id": me, "participants": [{"person_id": me}, {"person_id": bob}]})
    data = {"date": "2024-03-10", "direction": "expense", "amount": "10", "currency": "CNY", "content": "x"}
    row = service.update_txn(conn, t["id"], data, write_json=False)
    assert row["my_amount"] == "5.00"
    owed = sorted(p["owed"] for p in aa_service.get_txn_aa(conn, t["id"])["participants"])
    assert owed == ["5.00", "5.00"]


def test_update_amount_rejects_stale_fixed_split():
    _, conn, me, bob = _env()
    t = _txn(conn, "100", aa={"payer_id": me, "participants": [
        {"person_id": me, "fixed_amount": "30"}, {"person_id": bob, "fixed_amount": "70"}]})
    data = {"date": "2024-03-10", "direction": "expense", "amount": "10", "currency": "CNY", "content": "x"}
    with pytest.raises(ValueError):
        service.update_txn(conn, t["id"], data, write_json=False)
    assert service.get_txn(conn, t["id"])["amount"] == "100.00"   # rolled back
    # Sending the new split along with the amount works, atomically.
    data["aa"] = {"payer_id": me, "participants": [
        {"person_id": me, "fixed_amount": "3"}, {"person_id": bob, "fixed_amount": "7"}]}
    assert service.update_txn(conn, t["id"], data, write_json=False)["my_amount"] == "3.00"


def test_update_to_income_while_aa_is_rejected_and_null_aa_clears():
    _, conn, me, bob = _env()
    t = _txn(conn, "100", aa={"payer_id": me, "participants": [{"person_id": me}, {"person_id": bob}]})
    data = {"date": "2024-03-10", "direction": "income", "amount": "100", "currency": "CNY", "content": "x"}
    with pytest.raises(ValueError):
        service.update_txn(conn, t["id"], data, write_json=False)
    data["aa"] = None                         # explicit clear -> allowed
    row = service.update_txn(conn, t["id"], data, write_json=False)
    assert row["direction"] == "income" and not row["is_aa"]


def test_generate_without_payer_leaves_nothing():
    _, conn, me, bob = _env()
    s = aa_service.save_session(conn, {"name": "r", "currency": "CNY", "payer_id": None,
                                       "total_with_tax": "100",
                                       "items": [{"pretax": "100", "person_ids": [me, bob]}]},
                                write_json=False)
    with pytest.raises(ValueError):
        aa_service.generate_txn_from_session(conn, s["id"], date="2024-03-10", content="r",
                                             write_json=False)
    assert _count(conn) == 0


def test_generate_with_zero_pretax_leaves_nothing():
    _, conn, me, bob = _env()
    s = aa_service.save_session(conn, {"name": "r", "currency": "CNY", "payer_id": me,
                                       "total_with_tax": "50",
                                       "items": [{"pretax": "0", "person_ids": [me, bob]}]},
                                write_json=False)
    with pytest.raises(ValueError):
        aa_service.generate_txn_from_session(conn, s["id"], date="2024-03-10", content="r",
                                             write_json=False)
    assert _count(conn) == 0
    assert conn.execute("SELECT COUNT(*) FROM aa_session_txn").fetchone()[0] == 0


# --------------------------------------------------------------- validation
@pytest.mark.parametrize("bad", ["Infinity", "NaN", "1e30", "1e20"])
def test_non_finite_or_huge_amounts_are_value_errors(bad):
    with pytest.raises(ValueError):
        money.to_minor(bad, 2)
    with pytest.raises(ValueError):
        money.cny_to_base_micro(bad)


def test_infinite_percent_budget_rejected_and_dashboard_survives():
    _, conn, *_ = _env()
    for bad in ("Infinity", "NaN"):
        with pytest.raises(ValueError):
            budgets.set_budget(conn, category_id=_cat(conn, "饮食"), amount_type="percent",
                               percent=bad, write_json=False)
    dashboard.dashboard(conn)                 # no crash


def test_session_bad_input_is_value_error():
    _, conn, me, _ = _env()
    base = {"name": "r", "currency": "CNY", "payer_id": me, "total_with_tax": "10"}
    for items, total in (
        ([{"pretax": "1", "person_ids": [me]}], "-1"),
        ([{"pretax": "-1", "person_ids": [me]}], "10"),
        ([{"pretax": "1", "person_ids": [999]}], "10"),
        ([{"pretax": "1", "person_ids": [me, me]}], "10"),
        ([{"pretax": "1", "shares": [{"person_id": me, "weight": "-2"}]}], "10"),
    ):
        with pytest.raises(ValueError):
            aa_service.save_session(conn, {**base, "total_with_tax": total, "items": items},
                                    write_json=False)


# --------------------------------------------------------------- fiscal months
def _rate(conn, code, d, r):
    cid = conn.execute("SELECT id FROM currency WHERE code=?", (code,)).fetchone()["id"]
    conn.execute("INSERT INTO exchange_rate (currency_id, rate_date, rate_to_base) VALUES (?,?,?)",
                 (cid, d, str(r)))
    conn.commit()


def test_fiscal_usd_figures_ignore_card_currency():
    _, conn, *_ = _env()
    _rate(conn, "USD", "2024-03-01", "7.1")
    _rate(conn, "JPY", "2024-03-01", "0.05")
    _txn(conn, "710", date="2024-02-10")
    stats.save_fiscal_month(conn, "2024-03-05", credit_card="10000", credit_card_currency="JPY",
                            write_json=False)
    fm = stats.fiscal_month(conn, "2024-03-05")
    assert fm["usd_rate"] == 7.1 and fm["total_usd"] == 100.0
    assert fm["credit_card_cny"] == 500.0     # 10000 JPY * 0.05


def test_dashboard_honours_total_override():
    _, conn, *_ = _env()
    _txn(conn, "710", date="2024-08-10")      # -> 2024-09-05 month
    stats.save_fiscal_month(conn, "2024-09-05", credit_card="0", total_cny="999", write_json=False)
    d = dashboard.dashboard(conn, today="2024-09-20")
    assert d["previous"]["total_cny"] == 999.0


# --------------------------------------------------------------- xlsx sync
def _sync_env():
    d, conn, me, bob = _env()
    cny = conn.execute("SELECT id FROM currency WHERE code='CNY'").fetchone()["id"]
    food = _cat(conn, "饮食")
    for day, minor, content, note in (("2024-01-10", 1000, "咖啡", None),
                                      ("2024-01-11", 2550, "午饭", "with friend")):
        conn.execute(
            """INSERT INTO txn (external_uid, txn_date, direction, amount_minor, currency_id,
                   content, category_id, note, rate_used, base_amount_micro, revision)
               VALUES (?, ?, 'expense', ?, ?, ?, ?, ?, '1', ?, 1)""",
            (uuid.uuid4().hex, day, minor, cny, content, food, note, minor * 10000))
    conn.commit()
    xlsx = d / "book.xlsx"
    xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)
    return conn, xlsx


def _edit(xlsx, fn):
    wb = openpyxl.load_workbook(str(xlsx))
    fn(wb["支出"], wb)
    wb.save(str(xlsx)); wb.close()


def _diff(conn, xlsx):
    return xlsxsync.diff_from_xlsx(conn, xlsx)


def test_note_and_whitespace_do_not_cause_false_changes():
    conn, xlsx = _sync_env()
    assert _diff(conn, xlsx).summary()["unchanged"] == 2      # txn with a note included
    t = _txn(conn, "5", content="茶 ")                        # trailing space
    xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)
    s = _diff(conn, xlsx).summary()
    assert s["unchanged"] == 3 and s["modified"] == 0 and t


def test_note_row_edit_in_excel_is_modified_not_conflict():
    conn, xlsx = _sync_env()
    _edit(xlsx, lambda ws, wb: setattr(ws.cell(row=3, column=4), "value", "午饭 改"))
    s = _diff(conn, xlsx).summary()
    assert s["modified"] == 1 and s["conflict"] == 0


def test_unknown_category_is_anomaly():
    conn, xlsx = _sync_env()
    _edit(xlsx, lambda ws, wb: setattr(ws.cell(row=2, column=5), "value", "宠物"))
    s = _diff(conn, xlsx).summary()
    assert s["anomalies"] == 1 and s["modified"] == 0


def test_bad_row_is_anomaly_not_deletion():
    conn, xlsx = _sync_env()
    _edit(xlsx, lambda ws, wb: setattr(ws.cell(row=2, column=3), "value", "人民bi"))
    s = _diff(conn, xlsx).summary()
    assert s["anomalies"] == 1 and s["deleted"] == 0


def test_copy_pasted_uid_becomes_new_row():
    conn, xlsx = _sync_env()

    def paste(ws, wb):
        for c in range(1, 9):
            ws.cell(row=4, column=c).value = ws.cell(row=3, column=c).value
        ws.cell(row=4, column=4).value = "晚饭"
        ws.cell(row=4, column=2).value = 40
    _edit(xlsx, paste)
    diff = _diff(conn, xlsx)
    s = diff.summary()
    assert s["new"] == 1 and s["modified"] == 0 and s["unchanged"] == 2
    xlsxsync.apply_diff(conn, xlsx, diff, write_json=False)
    contents = sorted(r[0] for r in conn.execute("SELECT content FROM txn"))
    assert contents == ["午饭", "咖啡", "晚饭"]


def test_second_import_does_not_duplicate():
    conn, xlsx = _sync_env()

    def add(ws, wb):
        ws.cell(row=4, column=1).value = date(2024, 1, 12)
        ws.cell(row=4, column=2).value = 5
        ws.cell(row=4, column=3).value = "人民币"
        ws.cell(row=4, column=4).value = "新行"
    _edit(xlsx, add)
    xlsxsync.apply_diff(conn, xlsx, _diff(conn, xlsx), write_json=False)
    assert _count(conn) == 3
    # uid was written back, so a second import sees nothing new.
    s = _diff(conn, xlsx).summary()
    assert s["new"] == 0 and s["unchanged"] == 3
    xlsxsync.apply_diff(conn, xlsx, _diff(conn, xlsx), write_json=False)
    assert _count(conn) == 3


def test_migrated_rows_without_uid_are_recognised():
    """Rows imported by the importer (uid never written to the sheet) are matched
    by their deterministic uid instead of colliding on insert."""
    conn, xlsx = _sync_env()
    _edit(xlsx, lambda ws, wb: [setattr(ws.cell(row=r, column=8), "value", None) for r in (2, 3)])
    from cashlog import importer
    for r, (minor, content) in ((2, (1000, "咖啡")), (3, (2550, "午饭"))):
        h = importer._row_hash("expense", f"2024-01-1{r - 2}", minor, "CNY", content, "饮食", "")
        conn.execute("UPDATE txn SET external_uid = ? WHERE content = ?",
                     (importer.make_external_uid("支出", r, h), content))
    conn.commit()
    diff = _diff(conn, xlsx)
    assert diff.summary()["new"] == 0 and diff.summary()["unchanged"] == 2
    xlsxsync.apply_diff(conn, xlsx, diff, write_json=False)
    assert _count(conn) == 2
    wb = openpyxl.load_workbook(str(xlsx)); uids = [wb["支出"].cell(row=r, column=8).value for r in (2, 3)]
    wb.close()
    assert all(uids)


def test_export_refuses_to_wipe_unimported_rows():
    conn, xlsx = _sync_env()
    _edit(xlsx, lambda ws, wb: [setattr(ws.cell(row=4, column=c), "value", v)
                               for c, v in ((1, date(2024, 1, 12)), (2, 5), (3, "人民币"), (4, "新行"))])
    with pytest.raises(xlsxsync.PendingSheetChangesError):
        xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)
    wb = openpyxl.load_workbook(str(xlsx)); assert wb["支出"].cell(row=4, column=4).value == "新行"
    wb.close()


def test_row_moved_between_sheets_is_anomaly():
    conn, xlsx = _sync_env()

    def move(ws, wb):
        inc = wb["入账"]
        for c in range(1, 9):
            inc.cell(row=2, column=c).value = ws.cell(row=2, column=c).value
            ws.cell(row=2, column=c).value = None
        inc.cell(row=2, column=5).value = None
    _edit(xlsx, move)
    s = _diff(conn, xlsx).summary()
    assert s["anomalies"] == 1 and s["modified"] == 0 and s["deleted"] == 0


def test_stats_sheet_lists_auto_filled_months():
    conn, xlsx = _sync_env()                  # txns in Jan 2024, no stored fiscal_month
    wb = openpyxl.load_workbook(str(xlsx))
    months = [wb["统计"].cell(row=r, column=1).value for r in range(2, wb["统计"].max_row + 1)]
    wb.close()
    assert len([m for m in months if m]) == len(stats.fiscal_months(conn))


def test_rate_refresh_reaches_back_for_weekend_txn():
    _, conn, *_ = _env()
    _rate(conn, "USD", "2024-01-01", "7")
    _txn(conn, "1", date="2024-01-14", currency="USD")       # a Sunday, stale
    usd = conn.execute("SELECT id FROM currency WHERE code='USD'").fetchone()["id"]
    assert rates._default_refresh_start(conn, usd) <= "2024-01-12"
