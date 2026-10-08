"""Structured AA counts only my share (txn_eff view) — ledger, stats, budgets, xlsx.

The stored txn keeps the full amount; only the *effective* amount changes once
AA is structured — settled or not. Isolated temp DB / workbook; real data is never touched.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from cashlog import aa_service, budgets, db, seed, service, stats, xlsxsync  # noqa: E402


def _env():
    d = Path(tempfile.mkdtemp(prefix="cashlog_aaeff_"))
    db.init_db(d / "t.db", reset=True)
    conn = db.connect(d / "t.db")
    seed.seed(conn)
    me = conn.execute("SELECT id FROM person WHERE is_self = 1").fetchone()["id"]
    other = aa_service.create_person(conn, "LYH", write_json=False)["id"]
    cat = conn.execute("SELECT id FROM category WHERE name='饮食' AND kind='expense'").fetchone()["id"]
    return d, conn, me, other, cat


def _aa_txn(conn, me, other, cat, *, total="100", mine="30", payer=None, settle=True):
    t = service.create_txn(conn, {"date": "2024-03-10", "direction": "expense", "amount": total,
                                  "currency": "CNY", "content": "吃饭 AA", "category_id": cat},
                           write_json=False)["id"]
    rest = str(int(total) - int(mine))
    aa_service.set_txn_aa(conn, t, payer_id=payer or me, participants=[
        {"person_id": me, "fixed_amount": mine}, {"person_id": other, "fixed_amount": rest}],
        write_json=False)
    if settle:
        aa_service.settle(conn, [t], write_json=False)
    return t


def test_my_share_counts_before_and_after_settle():
    _, conn, me, other, cat = _env()
    t = _aa_txn(conn, me, other, cat, settle=False)
    row = service.get_txn(conn, t)
    # Unsettled already counts only my share, in the ledger and in stats.
    assert row["amount"] == "100.00" and row["my_amount"] == "30.00"
    assert stats.expense_total_micro(conn, "2024-03-05", "2024-04-05") == 30 * 1_000_000

    aa_service.settle(conn, [t], write_json=False)
    row = service.get_txn(conn, t)
    assert row["amount"] == "100.00"          # stored total untouched (edit form)
    assert row["my_amount"] == "30.00"
    assert row["my_base_cny"] == "30.00"
    assert stats.expense_total_micro(conn, "2024-03-05", "2024-04-05") == 30 * 1_000_000
    # The AA split is still derived from the full amount.
    owed = sorted(p["owed"] for p in aa_service.get_txn_aa(conn, t)["participants"])
    assert owed == ["30.00", "70.00"]


def test_stats_budgets_dashboard_count_my_share():
    _, conn, me, other, cat = _env()
    _aa_txn(conn, me, other, cat)                       # 100 total, my share 30
    service.create_txn(conn, {"date": "2024-03-11", "direction": "expense", "amount": "5",
                              "currency": "CNY", "content": "咖啡", "category_id": cat},
                       write_json=False)
    total = stats.expense_total_micro(conn, "2024-03-05", "2024-04-05")
    assert total == 35 * 1_000_000
    assert stats.category_total_micro(conn, cat, "2024-03-05", "2024-04-05") == 35 * 1_000_000
    bd = {r["category_id"]: r["total"] for r in
          stats.category_breakdown(conn, "2024-03-05", "2024-04-04")}
    assert bd[cat] == 35.0
    budgets.set_budget(conn, category_id=cat, amount_type="absolute", amount="50",
                       percent=None, note=None)
    st = budgets.budget_status(conn, period_end="2024-04-05")
    cat_row = next(c for c in st["categories"] if c["category_id"] == cat)
    assert float(cat_row["actual"]) == 35.0


def test_not_a_participant_counts_zero():
    """I fronted the whole bill for others: once settled, none of it is my spend."""
    _, conn, me, other, cat = _env()
    third = aa_service.create_person(conn, "DK", write_json=False)["id"]
    t = service.create_txn(conn, {"date": "2024-03-10", "direction": "expense", "amount": "40",
                                  "currency": "CNY", "content": "代付 AA", "category_id": cat},
                           write_json=False)["id"]
    aa_service.set_txn_aa(conn, t, payer_id=me, participants=[
        {"person_id": other}, {"person_id": third}], write_json=False)
    aa_service.settle(conn, [t], write_json=False)
    assert service.get_txn(conn, t)["my_amount"] == "0.00"
    assert stats.expense_total_micro(conn, "2024-03-05", "2024-04-05") == 0


def test_xlsx_round_trip_writes_my_share_and_keeps_total():
    d, conn, me, other, cat = _env()
    t = _aa_txn(conn, me, other, cat)
    xlsx = d / "book.xlsx"
    xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)

    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb["支出"]
    assert ws.cell(row=2, column=2).value == 30.0       # my share in the sheet
    wb.close()
    assert xlsxsync.diff_from_xlsx(conn, xlsx).summary()["unchanged"] == 1

    # A content edit in Excel applies, and the stored total survives.
    wb = openpyxl.load_workbook(str(xlsx))
    wb["支出"].cell(row=2, column=4, value="吃饭 AA 改")
    wb.save(str(xlsx)); wb.close()
    diff = xlsxsync.diff_from_xlsx(conn, xlsx)
    assert diff.summary()["modified"] == 1
    xlsxsync.apply_diff(conn, xlsx, diff, write_json=False)
    row = conn.execute("SELECT amount_minor, content FROM txn WHERE id=?", (t,)).fetchone()
    assert row["amount_minor"] == 10000 and row["content"] == "吃饭 AA 改"
    assert xlsxsync.diff_from_xlsx(conn, xlsx).summary()["unchanged"] == 1

    # An amount edit can't be mapped back to a total: reported, never applied.
    wb = openpyxl.load_workbook(str(xlsx))
    wb["支出"].cell(row=2, column=2, value=35.0)
    wb.save(str(xlsx)); wb.close()
    diff = xlsxsync.diff_from_xlsx(conn, xlsx)
    s = diff.summary()
    assert s["anomalies"] == 1 and s["modified"] == 0 and s["deleted"] == 0
    xlsxsync.apply_diff(conn, xlsx, diff, write_json=False)
    assert conn.execute("SELECT amount_minor FROM txn WHERE id=?", (t,)).fetchone()[0] == 10000


def test_structuring_aa_after_export_is_not_an_anomaly():
    """Structure an AA in the app after an export: the sheet still shows the old
    total, which is simply 'unchanged' (the next export rewrites it), not an error."""
    d, conn, me, other, cat = _env()
    t = service.create_txn(conn, {"date": "2024-03-10", "direction": "expense", "amount": "100",
                                  "currency": "CNY", "content": "吃饭 AA", "category_id": cat},
                           write_json=False)["id"]
    xlsx = d / "book.xlsx"
    xlsxsync.export_to_xlsx(conn, target=xlsx, backup=False, write_json=False)
    aa_service.set_txn_aa(conn, t, payer_id=me, participants=[
        {"person_id": me}, {"person_id": other}], write_json=False)
    s = xlsxsync.diff_from_xlsx(conn, xlsx).summary()
    assert s["unchanged"] == 1 and s["anomalies"] == 0 and s["modified"] == 0


def test_non_aa_has_no_my_amount():
    _, conn, *_ = _env()
    t = service.create_txn(conn, {"date": "2024-03-10", "direction": "expense", "amount": "8",
                                  "currency": "CNY", "content": "AA 文本但未结构化"},
                           write_json=False)["id"]
    assert service.get_txn(conn, t)["my_amount"] is None
